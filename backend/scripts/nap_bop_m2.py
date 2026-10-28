"""Nạp BOP-M2 (tình trạng vật liệu M001…M008) từ file Excel công trình của PC vào DB.

Nguồn: file .xlsx tên có mã công trình (CTxx-xxx) nằm trong thư mục "05_PRODUCTION STATUS", quét mọi thư mục con
  của thư mục gốc (mặc định G:\\Shared drives\\PPD-Production Planning Department\\005_PROJECT — cấu trúc
  NĂM 20xx\\<công trình>\\02_SẢN XUẤT\\05_PRODUCTION STATUS\\<tên>_CTxx-xxx.xlsx; các thư mục khác của công trình
  chứa hàng trăm file BOM / PR cũng mang mã CT nên phải lọc theo thư mục),
  sheet 02_DATA (1 dòng / HEX: 9 cột "BOP-M2 <nhóm vật tư>" + "TÌNH TRẠNG BOP-M2 FINAL")
  sheet 03_MATERIALS (1 dòng / vật liệu).
Đích: bảng bop_m2_hex, bop_m2_vat_lieu, bop_m2_file (tạo bằng backend/sql/2026-10-28_bop_m2.sql).

- Mỗi công trình (MÃ CÔNG TRÌNH ở 02_DATA, không có thì mã CTxx-xxx trong tên file) thay TOÀN BỘ dữ liệu cũ của nó
  trong 1 giao dịch. File lỗi (thiếu sheet / cột / không có HEX) => ghi lỗi vào bop_m2_file, dữ liệu cũ giữ nguyên.
- 1 công trình có nhiều file (bản sao, bản lưu trữ…) => chỉ lấy file SỬA GẦN NHẤT, in ra các file bị bỏ.
- Lối tắt (.lnk, vd "NĂM 2026" ở 001_PROJECT của BỘ PHẬN ĐIỀU PHỐI-PC) của Google Drive không đi theo được: chúng
  trỏ về 005_PROJECT ở trên. Thư mục khác => --thu-muc; tên thư mục chứa file khác => --thu-muc-file.
- Chỉ nạp lại file có mtime / kích thước khác lần trước (--tat-ca để nạp lại hết). Đọc kiểu read_only (từng dòng)
  nên chịu được file lớn (có file > 200 MB).
- Ngày ghi dạng chữ dd/mm/yyyy (giống giao diện).

Cách chạy (từ thư mục backend):
  python scripts/nap_bop_m2.py                      # nạp file mới / đã đổi
  python scripts/nap_bop_m2.py --tat-ca             # nạp lại mọi file
  python scripts/nap_bop_m2.py --thu-muc "<A>" --thu-muc "<B>"  # quét nhiều thư mục gốc
  python scripts/nap_bop_m2.py --file "<đường dẫn>"  # chỉ 1 file
  python scripts/nap_bop_m2.py --thu                # chỉ đọc file + in số, KHÔNG ghi DB
DATABASE_URL lấy từ biến môi trường, không có thì đọc backend/.env.
"""
import argparse, datetime, json, os, re, sys, time

import openpyxl
import warnings

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
sys.stdout.reconfigure(encoding='utf-8')

THU_MUC = r'G:\Shared drives\PPD-Production Planning Department\005_PROJECT'
# Chỉ lấy file nằm trong thư mục có tên này (so không phân biệt hoa thường)
THU_MUC_FILE = '05_PRODUCTION STATUS'
MA_CT_TEN = re.compile(r'CT\d{2}-\d{3}', re.I)
NHOM_VAT_TU = ['GỖ', 'TẤM VÁN', 'VẢI/DA/MÂY', 'KIM LOẠI', 'KÍNH', 'ĐÁ', 'PHỤ KIỆN', 'GIA CÔNG', 'KHÁC']
COT_VL = [  # (cột DB, tiêu đề sheet 03_MATERIALS)
    ('nhom_vat_tu', 'NHÓM VẬT TƯ'), ('code', 'CODE VẬT LIỆU'), ('mo_ta', 'MÔ TẢ VẬT LIỆU'), ('tt_mau', 'TÌNH TRẠNG MẪU'),
    ('ngay_duyet', 'NGÀY DUYỆT MẪU'), ('nhu_cau', 'PHÂN LOẠI NHU CẦU ĐẶT MUA'), ('ghi_chu', 'GHI CHÚ'),
    ('tt_pr', 'TÌNH TRẠNG PR'), ('ngay_pr', 'NGÀY PR'), ('ngay_du_kien', 'NGÀY DỰ KIẾN GIAO HÀNG PMH NHẬP'),
    ('ngay_ve', 'NGÀY THỰC TẾ VỀ'), ('bop_m2', 'BOP-M2'),
]


def chu(v):
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def chuan(v):
    return re.sub(r'\s+', ' ', chu(v)).upper()


def bang(ws, moc):
    """(lay, các dòng sau tiêu đề) — tiêu đề = dòng đầu (trong 10 dòng) có ô đúng bằng `moc`."""
    it = ws.iter_rows(values_only=True)
    for _ in range(10):
        hdr = next(it, None)
        if hdr is None:
            break
        if any(chuan(v) == moc for v in hdr):
            cot = {}
            for j, v in enumerate(hdr):
                k = chuan(v)
                if k and k not in cot:
                    cot[k] = j
            lay = lambda r, ten: chu(r[cot[ten]]) if ten in cot and cot[ten] < len(r) else ''
            return lay, cot, it
    raise ValueError(f'Sheet "{ws.title}" không có cột "{moc}"')


def doc_file(duong_dan):
    wb = openpyxl.load_workbook(duong_dan, read_only=True, data_only=True)
    try:
        if '02_DATA' not in wb.sheetnames:
            raise ValueError('Không có sheet 02_DATA')
        lay, cot, dong = bang(wb['02_DATA'], 'HEX')
        if 'TÌNH TRẠNG BOP-M2 FINAL' not in cot:
            raise ValueError('Sheet 02_DATA không có cột "TÌNH TRẠNG BOP-M2 FINAL" (mẫu file cũ)')
        ma_ct, hexs = '', {}
        for r in dong:
            h = lay(r, 'HEX')
            if not h:
                continue
            ma_ct = ma_ct or lay(r, 'MÃ CÔNG TRÌNH')
            m2 = {n: lay(r, f'BOP-M2 {n}') for n in NHOM_VAT_TU if lay(r, f'BOP-M2 {n}')}
            hexs[h] = dict(  # HEX trùng trong file: giữ dòng sau
                hex=h, ma_nha_may=lay(r, 'MÃ NHÀ MÁY'), ten_hang_muc=lay(r, 'TÊN HẠNG MỤC'),
                nhom_sp=lay(r, 'PHÂN LOẠI NHÓM SẢN PHẨM'), xuong=lay(r, 'XƯỞNG CHÍNH'),
                ipo=lay(r, 'TÌNH TRẠNG IPO FINAL'), bop=lay(r, 'BOP'),  # ipo = đúng cột lọc của pivot
                m2=m2, final=lay(r, 'TÌNH TRẠNG BOP-M2 FINAL'))
        if not hexs:
            raise ValueError('Sheet 02_DATA không có dòng HEX nào')
        vl = []
        if '03_MATERIALS' in wb.sheetnames:
            lay_m, _, dong_m = bang(wb['03_MATERIALS'], 'CODE VẬT LIỆU')
            for r in dong_m:
                d = {k: lay_m(r, t) for k, t in COT_VL}
                if d['code'] or d['mo_ta']:
                    d['stt'] = len(vl) + 1
                    vl.append(d)
        if not ma_ct:
            m = re.search(r'CT\d{2}-\d{3}', os.path.basename(duong_dan), re.I)
            ma_ct = m.group(0).upper() if m else ''
        if not ma_ct:
            raise ValueError('Không xác định được mã công trình')
        return ma_ct, list(hexs.values()), vl
    finally:
        wb.close()


def doc_env():
    url = os.environ.get('DATABASE_URL')
    if url:
        return url
    p = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), '.env')
    for dong in open(p, encoding='utf-8'):
        if dong.startswith('DATABASE_URL='):
            return dong.split('=', 1)[1].strip().strip('"')
    raise SystemExit('Không thấy DATABASE_URL')


def ket_noi():
    import psycopg2
    url = doc_env()
    cn = psycopg2.connect(url if 'sslmode=' in url else url + ('&' if '?' in url else '?') + 'sslmode=require')
    with cn.cursor() as c:
        c.execute('SET statement_timeout = 60000')
        c.execute("SET idle_in_transaction_session_timeout = 60000")
    cn.commit()
    return cn


def ghi(cn, ten_file, st, ma_ct, hexs, vl):
    from psycopg2.extras import execute_values
    with cn:  # 1 giao dịch: lỗi giữa chừng thì không mất dữ liệu cũ
        with cn.cursor() as c:
            c.execute('DELETE FROM bop_m2_hex WHERE ma_ct = %s', (ma_ct,))
            c.execute('DELETE FROM bop_m2_vat_lieu WHERE ma_ct = %s', (ma_ct,))
            execute_values(c, '''INSERT INTO bop_m2_hex (ma_ct, hex, ma_nha_may, ten_hang_muc, nhom_sp, xuong, ipo, bop, m2, final, ten_file)
                VALUES %s''', [(ma_ct, h['hex'], h['ma_nha_may'], h['ten_hang_muc'], h['nhom_sp'], h['xuong'], h['ipo'], h['bop'],
                                json.dumps(h['m2'], ensure_ascii=False), h['final'], ten_file) for h in hexs])
            if vl:
                cols = [k for k, _ in COT_VL]
                execute_values(c, f'INSERT INTO bop_m2_vat_lieu (ma_ct, stt, {", ".join(cols)}, ten_file) VALUES %s',
                               [(ma_ct, d['stt'], *[d[k] for k in cols], ten_file) for d in vl])
            ghi_file(c, ten_file, st, ma_ct, len(hexs), len(vl), None)


def ghi_file(c, ten_file, st, ma_ct, so_hex, so_vl, loi):
    c.execute('''INSERT INTO bop_m2_file (ten_file, ma_ct, mtime, kich_thuoc, so_hex, so_vat_lieu, nap_luc, loi)
        VALUES (%s, %s, to_timestamp(%s), %s, %s, %s, now(), %s)
        ON CONFLICT (ten_file) DO UPDATE SET ma_ct = EXCLUDED.ma_ct, mtime = EXCLUDED.mtime, kich_thuoc = EXCLUDED.kich_thuoc,
          so_hex = EXCLUDED.so_hex, so_vat_lieu = EXCLUDED.so_vat_lieu, nap_luc = now(), loi = EXCLUDED.loi''',
              (ten_file, ma_ct, st.st_mtime, st.st_size, so_hex, so_vl, loi))


def tim_file(goc, thu_muc_file=THU_MUC_FILE):
    """.xlsx / .xlsm tên có mã CTxx-xxx trong thư mục `thu_muc_file` dưới các thư mục gốc; 1 mã nhiều file => file sửa gần nhất."""
    theo_ma = {}
    for g in goc:
        for thu_muc, _, ten_files in os.walk(g):
            if os.path.basename(thu_muc).strip().upper() != thu_muc_file.upper():
                continue
            for t in ten_files:
                if t.startswith('~$') or not t.lower().endswith(('.xlsx', '.xlsm')):
                    continue
                m = MA_CT_TEN.search(t)
                if m:
                    f = os.path.join(thu_muc, t)
                    theo_ma.setdefault(m.group(0).upper(), []).append((os.stat(f).st_mtime, f))
    chon = []
    for ma, ds in sorted(theo_ma.items()):
        ds.sort(reverse=True)
        chon.append(ds[0][1])
        for _, f in ds[1:]:
            print(f'BỎ  {f} (trùng {ma}, lấy bản sửa gần nhất: {os.path.basename(ds[0][1])})')
    print(f'Tìm thấy {len(chon)} công trình ({sum(len(v) for v in theo_ma.values())} file)')
    return chon


def main():
    ap = argparse.ArgumentParser(description='Nạp BOP-M2 từ file Excel công trình vào DB')
    ap.add_argument('--thu-muc', action='append', help=f'thư mục gốc, quét cả thư mục con (dùng nhiều lần được); mặc định {THU_MUC}')
    ap.add_argument('--thu-muc-file', default=THU_MUC_FILE, help=f'chỉ lấy file trong thư mục tên này (mặc định "{THU_MUC_FILE}")')
    ap.add_argument('--file', action='append', help='chỉ nạp file này (dùng nhiều lần được)')
    ap.add_argument('--tat-ca', action='store_true', help='nạp lại cả file không đổi')
    ap.add_argument('--thu', action='store_true', help='chỉ đọc + in, không ghi DB')
    a = ap.parse_args()

    files = a.file or tim_file(a.thu_muc or [THU_MUC], a.thu_muc_file)
    cn = None if a.thu else ket_noi()
    da_nap = {}
    if cn:
        with cn.cursor() as c:
            c.execute('SELECT ten_file, extract(epoch FROM mtime), kich_thuoc FROM bop_m2_file')
            da_nap = {r[0]: (float(r[1] or 0), r[2]) for r in c.fetchall()}
        cn.commit()

    tong = {'nap': 0, 'bo_qua': 0, 'loi': 0}
    for f in files:
        ten, st = os.path.basename(f), os.stat(f)
        cu = da_nap.get(ten)
        if not a.tat_ca and cu and abs(cu[0] - st.st_mtime) < 1 and cu[1] == st.st_size:
            tong['bo_qua'] += 1
            continue
        t = time.time()
        try:
            ma_ct, hexs, vl = doc_file(f)
            if cn:
                ghi(cn, ten, st, ma_ct, hexs, vl)
            print(f'OK   {ten}: {ma_ct} · {len(hexs)} HEX · {len(vl)} vật liệu · {time.time() - t:.1f}s')
            tong['nap'] += 1
        except Exception as e:  # noqa: BLE001 — file lỗi không chặn file khác
            print(f'LỖI {ten}: {e}')
            tong['loi'] += 1
            if cn:
                cn.rollback()
                with cn, cn.cursor() as c:
                    ghi_file(c, ten, st, None, None, None, str(e)[:500])
    print(f"Xong: nạp {tong['nap']}, bỏ qua (không đổi) {tong['bo_qua']}, lỗi {tong['loi']}")
    if cn:
        cn.close()


if __name__ == '__main__':
    main()
